#!/usr/bin/env python3
"""Excel + site isimlerini Assos fotoğraf URL'lerine bağla."""

from __future__ import annotations

import json
import re
from pathlib import Path

from openpyxl import load_workbook

TR = str.maketrans("ğıüşöçıâîû", "giusociaiu")
STOP = {
    "ve",
    "ile",
    "mm",
    "cm",
    "kg",
    "ml",
    "gr",
    "tip",
    "tipi",
    "yeni",
    "serisi",
    "nolu",
    "numara",
    "set",
    "seti",
    "govde",
    "saten",
    "sari",
    "nikel",
    "krom",
    "siyah",
    "beyaz",
    "naturel",
    "paslanmaz",
    "gumus",
    "gri",
    "bne",
    "sn",
    "sm",
    "gn",
    "gm",
}


def norm(s: str) -> str:
    s = (s or "").lower().replace("i̇", "i").translate(TR)
    s = " ".join(re.sub(r"[^a-z0-9]+", " ", s).split())
    return s.replace("dormakaba", "dorma").replace("kapi kapatici", "hidrolik")


def tokens(s: str) -> set[str]:
    return {t for t in norm(s).split() if t not in STOP and len(t) > 1}


def codes(s: str) -> set[str]:
    n = norm(s)
    found = set()
    for raw in re.findall(r"[a-z]*\d+[a-z0-9]*", n.replace(" ", "")):
        if len(raw) >= 3:
            found.add(raw)
    for a, b in re.findall(r"\b([a-z]{1,5})\s+(\d{2,4}[a-z]?)\b", n):
        found.add(a + b)
    return found


def stem(s: str) -> str:
    n = norm(s)
    n = re.sub(
        r"\b(saten|sari|nikel|krom|siyah|beyaz|naturel|paslanmaz|gumus|gri|govde|barelli|barelsiz)\b",
        " ",
        n,
    )
    return " ".join(n.split())


def main() -> None:
    wc = json.loads(Path("/tmp/assos-import/wc-products.json").read_text())
    wc_items = []
    for p in wc:
        imgs = p.get("images") or []
        src = imgs[0].get("src") if imgs else None
        if not src:
            continue
        name = p.get("name") or ""
        wc_items.append(
            {
                "name": name,
                "norm": norm(name),
                "stem": stem(name),
                "tok": tokens(name),
                "codes": codes(name),
                "img": src,
            }
        )

    excel_names = []
    wb = load_workbook("/Users/ase/Downloads/STOK LİSTESİ (1).XLSX", data_only=True)
    ws = wb[wb.sheetnames[0]]
    for row in ws.iter_rows(min_row=3, values_only=True):
        turu, kodu, extra = row[2], row[3], row[4]
        if not turu or not kodu:
            continue
        if str(extra or "").strip() == str(turu).strip():
            continue
        excel_names.append(" ".join(str(kodu).split()))

    photo_map: dict[str, str] = {}

    def add(key: str, img: str) -> None:
        k = norm(key)
        if k and img:
            photo_map[k] = img

    for w in wc_items:
        add(w["name"], w["img"])
        add(w["stem"], w["img"])

    matched = 0
    for name in excel_names:
        n = norm(name)
        st = stem(name)
        c = codes(name)
        t = tokens(name)
        if n in photo_map or st in photo_map:
            photo_map[n] = photo_map.get(n) or photo_map[st]
            matched += 1
            continue

        best = None
        best_score = 0.0
        for w in wc_items:
            score = 0.0
            shared_codes = c & w["codes"]
            if shared_codes:
                score += 2.0 * len(shared_codes)
            shared_tok = t & w["tok"]
            if shared_tok:
                score += len(shared_tok) / max(len(t | w["tok"]), 1)
            if st and st == w["stem"]:
                score += 3
            if score > best_score:
                best_score = score
                best = w

        # En az bir model kodu veya güçlü isim örtüşmesi
        if best and (
            best_score >= 2.3
            or (best_score >= 1.15 and (c & best["codes"]))
            or (best["stem"] == st and len(st.split()) >= 3)
        ):
            add(name, best["img"])
            add(st, best["img"])
            matched += 1

    Path("scripts/photo-map.json").write_text(
        json.dumps(
            {
                "map": photo_map,
                "wc": [
                    {
                        "norm": w["norm"],
                        "stem": w["stem"],
                        "tok": sorted(w["tok"]),
                        "codes": sorted(w["codes"]),
                        "img": w["img"],
                    }
                    for w in wc_items
                ],
            },
            ensure_ascii=False,
        )
    )
    print("map keys", len(photo_map), "excel matched", matched, "wc", len(wc_items))


if __name__ == "__main__":
    main()
